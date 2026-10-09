"""
RAG (Retrieval-Augmented Generation) Service for AI Chat Agent
Handles document upload, vector storage, and semantic retrieval.
"""

import os
import uuid
import json
import threading
import io
from datetime import datetime
from typing import List, Dict, Optional, Tuple
from pathlib import Path

# RAG imports
import chromadb
from chromadb.config import Settings
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_openai import OpenAIEmbeddings

# Local imports
from config_loader import load_config, save_config

# Configuration
CONFIG_FILE = "rag_config.json"
RAG_LOCK = threading.Lock()


class RAGService:
    """Handles document upload, storage, and retrieval for RAG."""
    
    def __init__(self):
        self.config = self._load_config()
        self.client = None
        self.vector_store = None
        self.embeddings = None
        self.text_splitter = None
        self._initialize()
    
    def _load_config(self) -> dict:
        """Load RAG configuration from file."""
        if os.path.exists(CONFIG_FILE):
            try:
                with open(CONFIG_FILE, "r") as f:
                    return json.load(f)
            except (json.JSONDecodeError, IOError):
                pass
        return self._get_defaults()
    
    def _get_defaults(self) -> dict:
        """Get default RAG configuration."""
        return {
            "enabled": True,
            "url": os.getenv("CHROMA_DB_URL", "http://chromadb:8000"),
            "apiKey": os.getenv("CHROMA_API_KEY", ""),
            "collectionName": os.getenv("CHROMA_COLLECTION", "documents"),
            "topK": int(os.getenv("RAG_TOP_K", "4")),
            "similarityThreshold": float(os.getenv("RAG_THRESHOLD", "0.25")),
            "embeddingModel": os.getenv("EMBEDDING_MODEL", "text-embedding-3-small"),
            "chunkSize": int(os.getenv("CHUNK_SIZE", "500")),
            "chunkOverlap": int(os.getenv("CHUNK_OVERLAP", "50")),
            "documents": []
        }
    
    def _initialize(self):
        """Initialize RAG components."""
        try:
            # Initialize embeddings (using OpenAI API)
            api_key = os.getenv("LLM_API_KEY", "")
            if api_key:
                self.embeddings = OpenAIEmbeddings(
                    model=self.config.get("embeddingModel", "text-embedding-3-small"),
                    api_key=api_key
                )
            else:
                # Fallback to local embeddings if no API key
                from langchain_huggingface import HuggingFaceEmbeddings
                self.embeddings = HuggingFaceEmbeddings(
                    model_name="sentence-transformers/all-MiniLM-L6-v2"
                )
            
            # Initialize text splitter
            self.text_splitter = RecursiveCharacterTextSplitter(
                chunk_size=self.config.get("chunkSize", 500),
                chunk_overlap=self.config.get("chunkOverlap", 50)
            )
            
            # Initialize ChromaDB client
            chroma_path = os.getenv("CHROMA_PATH", "./chroma_data")
            
            # Create persistent client
            self.client = chromadb.PersistentClient(path=chroma_path)
            
            # Get or create collection
            collection_name = self.config.get("collectionName", "documents")
            self.collection = self.client.get_or_create_collection(name=collection_name)
            
            # Initialize LangChain Chroma wrapper
            self.vector_store = Chroma(
                client=self.client,
                collection_name=collection_name,
                embedding_function=self.embeddings
            )
            
            self.config["ready"] = True
            
        except Exception as e:
            print(f"Error initializing RAG: {e}")
            self.config["ready"] = False
            self.config["error"] = str(e)
    
    def upload_document(self, file_path: str, file_content: str = None, file_name: str = None) -> dict:
        """
        Upload and process a document.
        """
        try:
            if file_content is None:
                if not os.path.exists(file_path):
                    return {"success": False, "error": f"File not found: {file_path}"}
                
                with open(file_path, "rb") as f:
                    file_content = f.read()
            
            documents = self._parse_document(file_content, file_path, file_name)
            
            if not documents:
                return {"success": False, "error": "No content could be extracted from document"}
            
            chunks = self.text_splitter.split_documents(documents)
            
            if not chunks:
                return {"success": False, "error": "No chunks could be created from document"}
            
            doc_id = str(uuid.uuid4())
            chunk_ids = [f"{doc_id}_{i}" for i in range(len(chunks))]
            
            texts = [chunk.page_content for chunk in chunks]
            metadatas = [chunk.metadata for chunk in chunks]
            
            self.vector_store.add_texts(texts=texts, ids=chunk_ids, metadatas=metadatas)
            
            doc_info = {
                "id": doc_id,
                "name": file_name or os.path.basename(file_path),
                "path": file_path,
                "chunkCount": len(chunks),
                "uploadDate": datetime.now().isoformat(),
                "size": len(file_content) if isinstance(file_content, str) else len(str(file_content))
            }
            
            with RAG_LOCK:
                self.config["documents"].append(doc_info)
                self._save_config()
            
            return {
                "success": True,
                "docId": doc_id,
                "name": doc_info["name"],
                "chunks": len(chunks),
                "message": f"Successfully uploaded and chunked document"
            }
            
        except Exception as e:
            return {"success": False, "error": str(e)}
    
    def _parse_document(self, content, file_path: str, file_name: str = None) -> List[Document]:
        """Parse document content based on file extension."""
        documents = []
        filename = file_name or os.path.basename(file_path).lower()
        
        try:
            if filename.endswith(".pdf"):
                documents = self._parse_pdf(content, filename)
            elif filename.endswith(".docx"):
                documents = self._parse_docx(content, filename)
            elif filename.endswith(".txt"):
                documents = self._parse_txt(content, filename)
            elif filename.endswith(".md"):
                documents = self._parse_md(content, filename)
            else:
                documents = self._parse_txt(content, filename)
        except Exception as e:
            documents = [Document(
                page_content=f"Error parsing document: {str(e)}",
                metadata={"source": filename, "error": str(e)}
            )]
        
        return documents
    
    def _parse_pdf(self, content, filename: str) -> List[Document]:
        """Parse PDF document."""
        try:
            import PyPDF2
            if isinstance(content, str):
                content = content.encode("utf-8", errors="ignore")
            
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(content))
            text = ""
            for page in pdf_reader.pages:
                text += page.extract_text() or "" + "\n"
            
            return [Document(
                page_content=text,
                metadata={"source": filename, "type": "pdf"}
            )]
        except Exception as e:
            return [Document(
                page_content=f"PDF parsing error: {str(e)}",
                metadata={"source": filename, "type": "pdf", "error": str(e)}
            )]
    
    def _parse_docx(self, content, filename: str) -> List[Document]:
        """Parse DOCX document."""
        try:
            from docx import Document as DocxDocument
            if isinstance(content, str):
                content = content.encode("utf-8", errors="ignore")
            
            doc = DocxDocument(io.BytesIO(content))
            text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
            
            return [Document(
                page_content=text,
                metadata={"source": filename, "type": "docx"}
            )]
        except Exception as e:
            return [Document(
                page_content=f"DOCX parsing error: {str(e)}",
                metadata={"source": filename, "type": "docx", "error": str(e)}
            )]
    
    def _parse_txt(self, content, filename: str) -> List[Document]:
        """Parse text file."""
        if isinstance(content, bytes):
            content = content.decode("utf-8", errors="ignore")
        
        return [Document(
            page_content=content,
            metadata={"source": filename, "type": "txt"}
        )]
    
    def _parse_md(self, content, filename: str) -> List[Document]:
        """Parse Markdown file."""
        if isinstance(content, bytes):
            content = content.decode("utf-8", errors="ignore")
        
        return [Document(
            page_content=content,
            metadata={"source": filename, "type": "markdown"}
        )]
    
    def retrieve_relevant_docs(self, query: str, top_k: int = None, threshold: float = None) -> List[dict]:
        """Retrieve relevant documents for a query."""
        try:
            if not self.config.get("enabled", True) or not self.config.get("ready", False):
                return []
            
            k = top_k or self.config.get("topK", 4)
            sim_threshold = threshold or self.config.get("similarityThreshold", 0.25)
            
            results = self.vector_store.similarity_search_with_score(query, k=k)
            
            relevant_docs = []
            for doc, score in results:
                similarity = 1 - score
                
                if similarity >= sim_threshold:
                    relevant_docs.append({
                        "content": doc.page_content,
                        "metadata": doc.metadata,
                        "score": similarity,
                        "threshold": sim_threshold
                    })
            
            return relevant_docs
            
        except Exception as e:
            print(f"Error retrieving documents: {e}")
            return []
    
    def query_documents(self, query: str, context_window: int = 3) -> Tuple[str, List[dict]]:
        """Query documents and return augmented context."""
        relevant = self.retrieve_relevant_docs(query)
        
        if not relevant:
            return "No relevant documents found.", []
        
        context_parts = []
        docs_info = []
        
        for doc in relevant:
            source = doc["metadata"].get("source", "Unknown")
            chunk_info = f"[Source: {source}]\n{doc["content"][:500]}..."
            context_parts.append(chunk_info)
            docs_info.append(doc)
        
        context = "\n\n---\n\n".join(context_parts)
        augmented_prompt = f"""Based on the following document context:

{context}

User Question: {query}

Please answer the question based on the provided document context. Cite specific sources when relevant.
"""
        
        return augmented_prompt, docs_info
    
    def list_documents(self) -> List[dict]:
        """List all uploaded documents."""
        return self.config.get("documents", [])
    
    def delete_document(self, doc_id: str) -> dict:
        """Delete a document and its chunks from the vector store."""
        try:
            with RAG_LOCK:
                doc_index = None
                for i, doc in enumerate(self.config["documents"]):
                    if doc["id"] == doc_id:
                        doc_index = i
                        break
                
                if doc_index is None:
                    return {"success": False, "error": f"Document not found: {doc_id}"}
                
                chunk_ids = [f"{doc_id}_{i}" for i in range(self.config["documents"][doc_index]["chunkCount"])]
                
                try:
                    self.collection.delete(ids=chunk_ids)
                except Exception as e:
                    print(f"Warning: Could not delete from collection: {e}")
                
                deleted_doc = self.config["documents"].pop(doc_index)
                self._save_config()
                
                return {
                    "success": True,
                    "message": f"Deleted document: {deleted_doc["name"]}",
                    "deletedChunks": deleted_doc["chunkCount"]
                }
                
        except Exception as e:
            return {"success": False, "error": str(e)}
    
    def clear_all(self) -> dict:
        """Clear all documents from the vector store."""
        try:
            with RAG_LOCK:
                self.collection.delete(where={})
                self.config["documents"] = []
                self._save_config()
                
                return {
                    "success": True,
                    "message": "All documents cleared"
                }
        except Exception as e:
            return {"success": False, "error": str(e)}
    
    def get_stats(self) -> dict:
        """Get RAG statistics."""
        docs = self.config.get("documents", [])
        total_chunks = sum(doc.get("chunkCount", 0) for doc in docs)
        
        return {
            "documents": len(docs),
            "totalChunks": total_chunks,
            "enabled": self.config.get("enabled", True),
            "ready": self.config.get("ready", False),
            "collection": self.config.get("collectionName", "documents")
        }
    
    def _save_config(self):
        """Save configuration to file."""
        try:
            with open(CONFIG_FILE, "w") as f:
                json.dump(self.config, f, indent=4)
        except Exception as e:
            print(f"Error saving config: {e}")


rag_service = None


def get_rag_service() -> RAGService:
    """Get or create RAG service instance."""
    global rag_service
    if rag_service is None:
        rag_service = RAGService()
    return rag_service


def reload_rag_config():
    """Reload RAG configuration."""
    global rag_service
    rag_service = None
    return get_rag_service()
